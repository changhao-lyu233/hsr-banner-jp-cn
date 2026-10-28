# 第 5 步：把所有人工编码合并成一张分析表 data/analysis_phases.csv（一行 = 一期，共 65 期）
#
# 读入：
#   data/phases.csv                     分期表（前面三个脚本生成）
#   data/manual/character_coding.csv    强度、主推、SP、同位体、性别等
#   data/manual/story_scoring.xlsx      剧情打分 → 算铺垫分、当期分
#   data/manual/va_final.csv            声优档位
#   data/manual/buffs.csv               角色加强
#   data/raw/wiki_characters_raw.tsv    汉语 CV（character_coding 里没有）
# 输出：
#   data/analysis_phases.csv
#
# 这一步只做"合并和编码"，不碰名次。名次指标在下一步 compute_metrics.py 里算。
# 运行：在项目文件夹里输入  py scripts\build_analysis_table.py
#
# 输出各列的意思（1 = 是，0 = 否）：
#   series 大版本系列（1.x–4.x，控制变量用）
#   phase_no 第几期 | phase_kind 期的类型 | in_main 是否进主分析（51 期）| window_short 开池后不满 7 天就到数据截止日
#   promo 促销期 | collision 和绝区零或原神同一天开池 | buffed_rerun 本期里"加强后首次复刻"的角色
#   strength 强度 0/1/2 | T0 | main_dps 主C | promoted 主推 | sp SP 形态 | isomorph 同位体 | male 男性
#   type_group 按类型分组：上半T0 / 上半非T0 / 下半 / SP（SP 不分上下半，单独一组）
#   story_pre_max 铺垫分 | story_cur_max 当期分 | story_pre_n2 UP 前 2 分以上的行数（累计口径）
#   story_pre_decay 远近口径的铺垫分 | story_pre_gap 铺垫最高分距 UP 隔几个版本 | story_pre_rows UP 前一共几行
#   jp_va_tier / cn_va_tier 日配、国配档位 A/B/C | jp_va_A / cn_va_A 是否 A 档 | jp_va_idol 日配是否偶像/乐队企划成员

import csv
import re
from datetime import date
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
MANUAL = DATA / "manual"

# ---------- 研究设计里定下的几条名单（改设计时改这里） ----------

# 赠送本体的首次 UP：主分析剔除，稳健性里放回
FREE_COPY = {"真理医生", "丹恒•腾荒"}

# 促销（首充双倍重置或累充）开始的日子，都正好是某期的开池日
PROMO_DATES = {"2024-03-27", "2025-04-09", "2025-11-05", "2026-04-22"}

# 分析数据截止日（10-25 及之后只是导出时刻的快照）
DATA_END = date(2026, 10, 24)

# 剧情"远近口径"：最高分那一行距 UP 隔几个版本以上就减 1 分
DECAY_GAP = 10

# 同一个人的不同形态：打分表里写的是本体名字，SP 形态要用本体的行
PERSON = {
    "丹恒•饮月": "丹恒", "丹恒•腾荒": "丹恒", "大黑塔": "黑塔", "长夜月": "三月七",
    "忘归人": "停云", "姬子•启行": "姬子", "火花": "花火", "知更鸟•晴歌": "知更鸟",
    "砂金•戏浪": "砂金", "银狼LV.999": "银狼", "千冶•刃": "刃",
}

# wiki 上的日配名字和声优表里的写法不一样（简体字、旧字体），在这里对齐
JP_CV_ALIAS = {
    "三木真一郎": "三木眞一郎", "石见舞菜香": "石見舞菜香", "小仓唯": "小倉唯",
    "伊东健人": "伊東健人", "佐佐木未来": "佐々木未来", "花泽香菜": "花澤香菜",
    "上田丽奈": "上田麗奈", "佐仓绫音": "佐倉綾音", "田中理惠": "田中理恵",
    "关智一": "関智一",
}

# 版本顺序（用来判断"UP 之前"和"隔了几个版本"）
VERSIONS = ([f"1.{i}" for i in range(7)] + [f"2.{i}" for i in range(8)]
            + [f"3.{i}" for i in range(9)] + [f"4.{i}" for i in range(8)])
V_INDEX = {v: i for i, v in enumerate(VERSIONS)}


# ---------- 读文件的小工具 ----------

def read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def col(row, prefix):
    """按列名开头取值（有些列名很长，只写开头几个字）"""
    for k, v in row.items():
        if k.startswith(prefix):
            return v
    raise KeyError(prefix)


def main_cv(name):
    """一个角色有两个声优时（藿藿和尾巴、流萤和萨姆、知更鸟的歌手），只取第一个，也就是角色本人"""
    return re.split(r"&| / ", name)[0].strip()


def parse_date(text):
    """把 2025/9/2 或 2025-9-2 统一成 2025-09-02"""
    y, m, d = re.split(r"[/-]", text.strip())
    return date(int(y), int(m), int(d)).isoformat()


# ---------- 1. 读入各张表 ----------

phases = read_csv(DATA / "phases.csv")
coding = {r["角色"]: r for r in read_csv(MANUAL / "character_coding.csv")}

va = {}
for r in read_csv(MANUAL / "va_final.csv"):
    va[(r["语言"], r["声优"])] = r

wiki_lines = (ROOT / "data/raw/wiki_characters_raw.tsv").read_text(encoding="utf-8").splitlines()
wiki = {r["角色"]: r for r in csv.DictReader([l for l in wiki_lines if not l.startswith("#")], delimiter="\t")}

buffs = read_csv(MANUAL / "buffs.csv")
buff_rerun_by_date = {}          # 开池日期 → 这一期里"加强后首次复刻"的角色
for b in buffs:
    d = col(b, "首次复刻开池日期").strip()
    if d:
        buff_rerun_by_date.setdefault(parse_date(d), []).append(b["角色"])

# 剧情打分：两页合在一起；"其他线索"页没填分的按 1 分
wb = load_workbook(MANUAL / "story_scoring.xlsx", read_only=True)
story_rows = []
for r in wb["重点剧情打分"].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    story_rows.append({"version": str(r[0]), "story": r[1], "person": r[2],
                       "score": int(r[8]), "sheet": "重点"})
for r in wb["其他线索"].iter_rows(min_row=2, values_only=True):
    if r[0] is None:
        continue
    story_rows.append({"version": str(r[0]), "story": r[2], "person": r[3],
                       "score": int(r[9]) if r[9] is not None else 1, "sheet": "其他"})
bad = {s["version"] for s in story_rows if s["version"] not in V_INDEX}
assert not bad, f"打分表里有认不出的版本号：{bad}"


# ---------- 2. 剧情变量 ----------

def story_vars(name, up_version):
    """按 data/manual/README.md 里定的规则，算一个角色的铺垫分、当期分等"""
    person = PERSON.get(name, name)
    up = V_INDEX[up_version]
    mine = [s for s in story_rows if s["person"] == person]
    before = [s for s in mine if V_INDEX[s["version"]] < up]
    same = [s for s in mine if V_INDEX[s["version"]] == up]
    # 例外：1.2 的「陌生女人的来信」在 1.2 下半才上线，对 1.2 上半的刃不算
    if name == "刃":
        same = [s for s in same if "陌生女人的来信" not in s["story"]]

    pre_max = max((s["score"] for s in before), default=0)
    # 最高分那一行离 UP 隔了几个版本（同分取最近的那行）
    if before:
        last_best = max(V_INDEX[s["version"]] for s in before if s["score"] == pre_max)
        gap = up - last_best
    else:
        gap = ""
    pre_decay = pre_max - 1 if (gap != "" and gap >= DECAY_GAP and pre_max > 0) else pre_max
    return {
        "story_pre_max": pre_max,                                    # 主分析：铺垫分
        "story_cur_max": max((s["score"] for s in same), default=0),  # 主分析：当期分
        "story_pre_n2": sum(s["score"] >= 2 for s in before),          # 稳健性：累计口径
        "story_pre_decay": pre_decay,                                  # 稳健性：远近口径
        "story_pre_gap": gap,
        "story_pre_rows": len(before),
    }


# ---------- 3. 声优档位 ----------

unmatched_cv = []

def va_tier(lang, cv_name, char):
    key = (lang, cv_name)
    if key not in va:
        unmatched_cv.append(f"{char}：{lang} {cv_name}")
        return {"tier": "", "A": "", "idol": ""}
    r = va[key]
    return {"tier": col(r, "最终档位"), "A": col(r, "主分析编码"),
            "idol": col(r, "日配：偶像") if lang == "日配" else ""}


# ---------- 4. 逐期合并 ----------

starts = [p["start_date"] for p in phases]
out = []
for i, p in enumerate(phases):
    d = p["start_date"]
    new = p["new_five"].strip()
    row = {
        "phase_no": i + 1,
        "start_date": d,
        "year": d[:4],
        "series": p["version"][0] + ".x",      # 大版本系列：星铁按"一年一个星球"策划，控制变量用它
        "version": p["version"],
        "half": p["half"],
        "new_five": new,
        "rerun_five": p["rerun_five"],
        "n_banners": p["n_banners"],
    }

    # 期的类型和是否进主分析
    if i == 0:
        kind = "开服首期"
    elif p["has_collab"] == "1":
        kind = "Fate联动"
    elif not new:
        kind = "纯复刻"
    elif new in FREE_COPY:
        kind = "新角色(赠送本体)"
    else:
        kind = "新角色"
    row["phase_kind"] = kind
    row["in_main"] = int(kind == "新角色")

    # 窗口是否完整：开池后不到 7 天就到数据截止日
    row["window_short"] = int((DATA_END - date.fromisoformat(d)).days < 6)

    # 期层面的标记
    row["promo"] = int(d in PROMO_DATES)
    row["zzz_same_day"] = p["zzz_same_day"]
    row["gi_same_day"] = p["gi_same_day"]
    row["collision"] = int(p["zzz_same_day"] == "1" or p["gi_same_day"] == "1")
    row["has_jilu"] = p["has_jilu"]
    row["has_long_banner"] = p["has_long_banner"]
    row["carryover_five"] = p["carryover_five"]
    row["buffed_rerun"] = "、".join(buff_rerun_by_date.get(d, []))

    # 新五星的属性（联动期有两个新角色，不进主分析，只留名字）
    if kind in ("新角色", "新角色(赠送本体)", "开服首期"):
        c = coding[new]
        up_version = re.match(r"\d\.\d", c["首次UP版本"]).group(0)
        strength = col(c, "强度")
        role = col(c, "定位")
        sp = c["SP形态(1/0)"]
        row.update({
            "strength": strength,
            "T0": int(strength == "0"),
            "main_dps": int(role in ("C", "主C/副C")),
            "role": role,
            "promoted": c["是否主推(1/0)"],
            "sp": sp,
            "isomorph": col(c, "同位体("),
            "male": int(c["性别(自动)"] == "男"),
            "body": c["体型(自动)"],
            "faction": c["阵营(自动)"],
        })
        # 按类型分组（米哈游的排法）：SP 优先单独成组，其余按上半/下半和强度分
        if sp == "1":
            group = "SP"
        elif p["half"] == "上半" and strength == "0":
            group = "上半T0"
        elif p["half"] == "上半":
            group = "上半非T0"
        else:
            group = "下半"
        row["type_group"] = group

        row.update(story_vars(new, up_version))

        jp_cv = main_cv(c["日语CV(自动)"])
        jp_cv = JP_CV_ALIAS.get(jp_cv, jp_cv)
        cn_cv = main_cv(wiki[new]["汉语CV"])
        j = va_tier("日配", jp_cv, new)
        k = va_tier("国配", cn_cv, new)
        row.update({"jp_cv": jp_cv, "jp_va_tier": j["tier"], "jp_va_A": j["A"], "jp_va_idol": j["idol"],
                    "cn_cv": cn_cv, "cn_va_tier": k["tier"], "cn_va_A": k["A"]})
    out.append(row)


# ---------- 5. 写出 ----------

fields = []
for r in out:
    for k in r:
        if k not in fields:
            fields.append(k)
with open(DATA / "analysis_phases.csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fields, restval="")
    w.writeheader()
    w.writerows(out)


# ---------- 6. 打印核对 ----------

from collections import Counter
main = [r for r in out if r["in_main"]]
print(f"写出 data/analysis_phases.csv：{len(out)} 期")
print("期的类型：", dict(Counter(r["phase_kind"] for r in out)))
print(f"主分析样本：{len(main)} 期（应为 51）")
print("按类型分组：", dict(Counter(r["type_group"] for r in main)))
print("强度 T0/T1/T2：", dict(Counter(r["strength"] for r in main)))
print("主推：", dict(Counter(r["promoted"] for r in main)), " SP：", dict(Counter(r["sp"] for r in main)),
      " 同位体：", dict(Counter(r["isomorph"] for r in main)), " 男性：", sum(r["male"] for r in main))
print("铺垫分 0/1/2/3：", dict(sorted(Counter(r["story_pre_max"] for r in main).items())))
print("当期分 0/1/2/3：", dict(sorted(Counter(r["story_cur_max"] for r in main).items())))
print("日配档位：", dict(Counter(r["jp_va_tier"] for r in main)), " 国配档位：", dict(Counter(r["cn_va_tier"] for r in main)))
print("促销期：", [r["start_date"] for r in out if r["promo"]])
print("撞车期：", [r["start_date"] for r in out if r["collision"]])
print("带加强复刻的期：", [(r["start_date"], r["buffed_rerun"]) for r in out if r["buffed_rerun"]])
print("窗口不满 7 天的期：", [r["start_date"] for r in out if r["window_short"]])
print("没对上声优表的：", unmatched_cv if unmatched_cv else "无")
